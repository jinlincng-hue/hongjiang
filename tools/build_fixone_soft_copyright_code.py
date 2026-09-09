from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT / "fixone-cloud-work"
OUT_DIR = ROOT / "outputs" / "soft-copyright"
DOCX_OUT = OUT_DIR / "Fixone视修工坊软件著作权源代码60页.docx"
TXT_OUT = OUT_DIR / "Fixone视修工坊软件著作权源代码60页.txt"
MANIFEST_OUT = OUT_DIR / "Fixone视修工坊软著代码选取说明.txt"

LINES_PER_PAGE = 50
TOTAL_PAGES = 60
TOTAL_LINES = LINES_PER_PAGE * TOTAL_PAGES

SELECTED_FILES = [
    "server/server.js",
    "server/db.js",
    "script.js",
    "style.css",
    "index.html",
    "profile.html",
    "upload-project.html",
    "repair-player.html",
    "admin-projects.html",
    "data/repair_flows.json",
    "data/products_demo.json",
    "data/users.js",
    "data/repair_clips.json",
    "seed.js",
    "package.json",
]


def read_lines(path: Path):
    text = path.read_text(encoding="utf-8", errors="replace")
    return text.splitlines()


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_run_font(run, name="Consolas", size=Pt(7.5)):
    run.font.name = name
    run.font.size = size
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), "DengXian")


def add_code_line(doc, number, file_name, source_line_number, content):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.left_indent = Inches(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    prefix = f"{number:04d}  {file_name}:{source_line_number:<4}  "
    visible_content = content.expandtabs(2)
    run = paragraph.add_run(prefix + visible_content)
    set_run_font(run)


def add_page_header(doc, page_number, first_file, last_file):
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    widths = [Inches(2.1), Inches(3.4), Inches(1.1)]
    labels = [
        "Fixone视修工坊源代码",
        f"{first_file} 至 {last_file}",
        f"第 {page_number:02d} 页",
    ]
    for idx, cell in enumerate(table.rows[0].cells):
        cell.width = widths[idx]
        set_cell_shading(cell, "F2F2F2")
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 2 else WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(labels[idx])
        set_run_font(run, "DengXian", Pt(7))
        run.bold = True


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    records = []
    source_summary = []
    for rel in SELECTED_FILES:
        path = PROJECT / rel
        if not path.exists():
            continue
        lines = read_lines(path)
        source_summary.append((rel, len(lines)))
        for idx, line in enumerate(lines, start=1):
            records.append((rel.replace("\\", "/"), idx, line))

    if len(records) < TOTAL_LINES:
        raise RuntimeError(f"selected source has only {len(records)} lines, need {TOTAL_LINES}")

    selected = records[:TOTAL_LINES]

    txt_lines = []
    for index, (rel, line_no, text) in enumerate(selected, start=1):
        page = (index - 1) // LINES_PER_PAGE + 1
        if (index - 1) % LINES_PER_PAGE == 0:
            txt_lines.append(f"\f第 {page:02d} 页  Fixone视修工坊源代码")
        txt_lines.append(f"{index:04d}  {rel}:{line_no:<4}  {text}")
    TXT_OUT.write_text("\n".join(txt_lines) + "\n", encoding="utf-8")

    manifest = [
        "Fixone视修工坊软件著作权源代码选取说明",
        "",
        "本材料仅选取 fixone-cloud-work 目录下的 Fixone 视修工坊自有源码。",
        "未选取 node_modules、dist、压缩包、构建产物、数据库文件以及红匠助修相关源码。",
        f"分页规则：共 {TOTAL_PAGES} 页，每页 {LINES_PER_PAGE} 行，合计 {TOTAL_LINES} 行。",
        "",
        "选取文件及源文件行数：",
    ]
    manifest.extend([f"- {rel}: {count} 行" for rel, count in source_summary])
    MANIFEST_OUT.write_text("\n".join(manifest) + "\n", encoding="utf-8")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)

    styles = doc.styles
    styles["Normal"].font.name = "Consolas"
    styles["Normal"].font.size = Pt(7.5)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "DengXian")

    for page in range(1, TOTAL_PAGES + 1):
        start = (page - 1) * LINES_PER_PAGE
        chunk = selected[start:start + LINES_PER_PAGE]
        add_page_header(doc, page, chunk[0][0], chunk[-1][0])
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(1)
        spacer.paragraph_format.space_before = Pt(0)
        for offset, (rel, line_no, text) in enumerate(chunk, start=start + 1):
            add_code_line(doc, offset, rel, line_no, text)
        if page < TOTAL_PAGES:
            doc.add_page_break()

    doc.core_properties.title = "Fixone视修工坊软件著作权源代码60页"
    doc.core_properties.subject = "软件著作权源代码材料"
    doc.core_properties.author = "Fixone"
    doc.save(DOCX_OUT)
    print(DOCX_OUT)
    print(TXT_OUT)
    print(MANIFEST_OUT)


if __name__ == "__main__":
    main()
